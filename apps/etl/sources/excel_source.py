"""
Excel/XLSX Source - the real UCI Online Retail II dataset's actual
format. Streams rows via openpyxl's read_only mode rather than
pandas.read_excel, which loads a whole sheet into memory before
returning anything (measured at ~75s / ~249MB for the real 1.07M-row
workbook during Volume 4.1 inspection - not acceptable as the default
read strategy on the proposal's 8GB-RAM target hardware).

Defaults to reading every sheet in the workbook, in order, because the
real UCI file needs both of its sheets ("Year 2009-2010" and
"Year 2010-2011") ingested - restricting to one sheet would silently
drop half the real dataset.
"""

import logging
from pathlib import Path
from typing import Iterator, Optional, Sequence, Union

import openpyxl
import pandas as pd

from etl.exceptions import MalformedSourceError, MissingColumnsError, SourceNotFoundError
from etl.sources.contract import CANONICAL_COLUMNS, SourceBatch

logger = logging.getLogger(__name__)


class ExcelSource:
    """Usage: `with ExcelSource(path) as source: for batch in source.extract(): ...`
    Pass `sheet_names=[...]` to restrict to specific sheets instead of
    reading the whole workbook."""

    def __init__(
        self,
        path: Union[str, Path],
        sheet_names: Optional[Sequence[str]] = None,
        batch_size: int = 50_000,
    ):
        self.path = Path(path)
        self._requested_sheets = list(sheet_names) if sheet_names else None
        self.sheet_names: list = []
        self.batch_size = batch_size
        self._workbook = None

    def __enter__(self) -> "ExcelSource":
        if not self.path.exists():
            raise SourceNotFoundError(f"Excel workbook not found: {self.path}")
        try:
            self._workbook = openpyxl.load_workbook(self.path, read_only=True, data_only=True)
        except Exception as exc:
            # openpyxl raises several distinct exception types for corrupt/non-xlsx files
            raise MalformedSourceError(f"Could not open workbook {self.path}: {exc}") from exc

        available = self._workbook.sheetnames
        requested = self._requested_sheets or available
        unknown = set(requested) - set(available)
        if unknown:
            self._workbook.close()
            self._workbook = None
            raise MalformedSourceError(
                f"Unknown sheet(s) {sorted(unknown)} in {self.path}; available sheets: {available}"
            )
        self.sheet_names = requested
        return self

    def __exit__(self, exc_type, exc_val, exc_tb) -> None:
        if self._workbook is not None:
            self._workbook.close()
            self._workbook = None

    def extract(self) -> Iterator[SourceBatch]:
        if self._workbook is None:
            raise RuntimeError(
                "ExcelSource must be used as a context manager: `with ExcelSource(...) as s:`"
            )
        batch_number = 0
        for sheet_name in self.sheet_names:
            worksheet = self._workbook[sheet_name]
            rows_iter = worksheet.iter_rows(values_only=True)
            try:
                header = list(next(rows_iter))
            except StopIteration:
                logger.warning(
                    "excel_source.empty_sheet", extra={"path": str(self.path), "sheet": sheet_name}
                )
                continue

            missing = set(CANONICAL_COLUMNS) - set(header)
            if missing:
                raise MissingColumnsError(
                    f"{self.path} sheet '{sheet_name}' is missing columns: {sorted(missing)}"
                )

            buffer = []
            for row in rows_iter:
                buffer.append(row)
                if len(buffer) >= self.batch_size:
                    yield self._make_batch(header, buffer, sheet_name, batch_number)
                    batch_number += 1
                    buffer = []
            if buffer:
                yield self._make_batch(header, buffer, sheet_name, batch_number)
                batch_number += 1

    def _make_batch(self, header, rows, sheet_name: str, batch_number: int) -> SourceBatch:
        df = pd.DataFrame(rows, columns=header)
        logger.info(
            "excel_source.batch",
            extra={
                "path": str(self.path),
                "sheet": sheet_name,
                "batch": batch_number,
                "rows": len(df),
            },
        )
        return SourceBatch(
            data=df, source_name=f"xlsx:{self.path.name}:{sheet_name}", batch_number=batch_number
        )
